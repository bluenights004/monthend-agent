"""Build the management pack figures from a snapshot, then write them into the Excel template.

All arithmetic is Decimal and happens here, not in Excel and not in a model. The
template is filled by *keys*, not cell addresses: hidden column A holds row keys
(e.g. `pl.rev.contract`) and hidden row 3 holds column keys (e.g. `month_actual`).
A firm can restyle, reorder or add rows to its template without touching code, as
long as the keys stay (docs/07-customising-the-template.md).
"""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from monthend.config import ClientConfig, Period
from monthend.money import ZERO, dec, pct, q2
from monthend.snapshot import SourceBundle
from monthend.sources.xero_parse import tb_imbalance

COLUMNS = ("month_actual", "month_budget", "month_var", "month_var_pct", "prior_month",
           "ytd_actual", "ytd_budget", "ytd_var")
BS_COLUMNS = ("closing", "prior_closing", "movement")
CREDIT_NATURES = {"income", "liability", "equity"}


@dataclass
class PackLine:
    key: str
    label: str
    statement: str
    type: str
    nature: str
    values: dict[str, Decimal | None]
    accounts: list[str] = field(default_factory=list)
    favourable: bool | None = None
    material: bool = False
    materiality_reason: str = ""

    @property
    def variance(self) -> Decimal:
        return self.values.get("month_var") or ZERO


@dataclass
class Check:
    code: str
    severity: str          # "block" stops submission for review; "warn" is shown to the reviewer
    message: str
    detail: dict = field(default_factory=dict)


@dataclass
class Pack:
    client_id: str
    client_name: str
    currency: str
    period: str
    period_label: str
    lines: list[PackLine]
    checks: list[Check]
    account_map: dict[str, str]
    generated_at: str

    def line(self, key: str) -> PackLine:
        for line in self.lines:
            if line.key == key:
                return line
        raise KeyError(key)

    def material_lines(self) -> list[PackLine]:
        return [l for l in self.lines if l.material]

    @property
    def blocking(self) -> list[Check]:
        return [c for c in self.checks if c.severity == "block"]

    def to_json(self) -> dict:
        return json.loads(json.dumps(asdict(self), default=str))

    @classmethod
    def from_json(cls, data: dict) -> "Pack":
        lines = []
        for raw in data["lines"]:
            values = {k: (dec(v) if v is not None else None) for k, v in raw["values"].items()}
            lines.append(PackLine(**{**raw, "values": values}))
        checks = [Check(**c) for c in data["checks"]]
        return cls(**{**data, "lines": lines, "checks": checks})


# --------------------------------------------------------------------------- mapping

def _in_range(code: str, low: str, high: str) -> bool:
    if code.isdigit() and low.isdigit() and high.isdigit():
        return int(low) <= int(code) <= int(high)
    return False


def map_accounts(rows: list[dict], codes: list[str]) -> tuple[dict[str, str], list[Check]]:
    """Assign each account code to a pack line. Explicit lists win; ranges are fallbacks."""
    explicit = {code: row["key"] for row in rows if row["type"] == "line" for code in row.get("accounts", [])}
    mapping, checks = {}, []
    for code in codes:
        if code in explicit:
            mapping[code] = explicit[code]
            continue
        fallback = next((row["key"] for row in rows if row["type"] == "line"
                         for low, high in row.get("ranges", []) if _in_range(code, low, high)), None)
        if fallback:
            mapping[code] = fallback
            checks.append(Check("MAP_FALLBACK", "warn",
                                f"Account {code} is not in the mapping; placed in '{fallback}' by code range. Confirm or update the mapping.",
                                {"account": code, "line": fallback}))
        else:
            checks.append(Check("MAP_MISSING", "block", f"Account {code} is not mapped to any pack line.", {"account": code}))
    return mapping, checks


# --------------------------------------------------------------------------- build

def _display(nature: str, signed: Decimal) -> Decimal:
    """Debit-positive TB values to presentation sign (income/liabilities/equity shown positive)."""
    return -signed if nature in CREDIT_NATURES else signed


def _is_material(variance: Decimal, budget: Decimal, rules: dict) -> tuple[bool, str]:
    abs_threshold, floor, pct_threshold = dec(rules["abs_threshold"]), dec(rules["abs_floor"]), dec(rules["pct_threshold"])
    size = abs(variance)
    if size >= abs_threshold:
        return True, f"variance >= {abs_threshold:,.0f}"
    if size >= floor:
        if budget == 0:
            return True, f"variance >= {floor:,.0f} against a nil budget"
        if abs(variance) / abs(budget) * 100 >= pct_threshold:
            return True, f"variance >= {floor:,.0f} and >= {pct_threshold}% of budget"
    return False, ""


def build_pack(client: ClientConfig, period: Period, bundle: SourceBundle, mapping_rows: list[dict]) -> Pack:
    checks: list[Check] = []
    imbalance = tb_imbalance(bundle.tb_current)
    if imbalance != 0:
        checks.append(Check("TB_OUT_OF_BALANCE", "block", f"Trial balance does not balance (difference {imbalance:,.2f}).",
                            {"difference": str(imbalance)}))
    prior_imbalance = tb_imbalance(bundle.tb_prior)
    if prior_imbalance != 0:
        checks.append(Check("TB_PRIOR_OUT_OF_BALANCE", "block",
                            f"Prior-month trial balance does not balance (difference {prior_imbalance:,.2f}).", {}))

    codes = sorted({l.code for l in bundle.tb_current} | {l.code for l in bundle.tb_prior} | set(bundle.budgets))
    account_map, map_checks = map_accounts(mapping_rows, codes)
    checks.extend(map_checks)

    current = {l.code: l for l in bundle.tb_current}
    prior = {l.code: l for l in bundle.tb_prior}
    fy = [p.key for p in period.fy_months(client.fy_start_month)]
    rows_by_key = {row["key"]: row for row in mapping_rows}

    def budget_for(code: str, months: list[str]) -> Decimal:
        return sum((bundle.budgets.get(code, {}).get(m, ZERO) for m in months), ZERO)

    values: dict[str, dict[str, Decimal | None]] = {}
    lines: list[PackLine] = []
    for row in mapping_rows:
        key, nature = row["key"], row["nature"]
        members = sorted(code for code, line_key in account_map.items() if line_key == key)
        if row["type"] == "line" and row["statement"] == "pl":
            actual = _display(nature, sum((current[c].month for c in members if c in current), ZERO))
            prior_m = _display(nature, sum((prior[c].month for c in members if c in prior), ZERO))
            ytd = _display(nature, sum((current[c].ytd for c in members if c in current), ZERO))
            # Budgets are entered in natural sign (income and expense both positive).
            budget = sum((budget_for(c, [period.key]) for c in members), ZERO)
            ytd_budget = sum((budget_for(c, fy) for c in members), ZERO)
            vals = {"month_actual": actual, "month_budget": budget, "prior_month": prior_m,
                    "ytd_actual": ytd, "ytd_budget": ytd_budget}
        elif row["type"] == "line":
            closing = _display(nature, sum((current[c].ytd for c in members if c in current), ZERO))
            prior_c = _display(nature, sum((prior[c].ytd for c in members if c in prior), ZERO))
            vals = {"closing": closing, "prior_closing": prior_c}
        elif row["type"] == "earnings":
            pl_codes = [c for c, k in account_map.items() if rows_by_key[k]["statement"] == "pl"]
            vals = {"closing": -sum((current[c].ytd for c in pl_codes if c in current), ZERO),
                    "prior_closing": -sum((prior[c].ytd for c in pl_codes if c in prior), ZERO)}
        else:  # total
            vals = {}
            for term in row["sum"]:
                sign, ref = (-1, term[1:]) if term.startswith("-") else (1, term)
                for col, v in values[ref].items():
                    if col in ("month_actual", "month_budget", "prior_month", "ytd_actual", "ytd_budget", "closing", "prior_closing"):
                        vals[col] = vals.get(col, ZERO) + sign * (v or ZERO)
        if row["statement"] == "pl":
            vals["month_var"] = vals["month_actual"] - vals["month_budget"]
            vals["month_var_pct"] = pct(vals["month_var"], vals["month_budget"])
            vals["ytd_var"] = vals["ytd_actual"] - vals["ytd_budget"]
        else:
            vals["movement"] = vals["closing"] - vals["prior_closing"]
        values[key] = vals

        line = PackLine(key, row["label"], row["statement"], row["type"], nature, vals, members)
        if row["statement"] == "pl":
            income_like = nature == "income"
            line.favourable = None if vals["month_var"] == 0 else (vals["month_var"] > 0) == income_like
            if row["type"] == "line":
                line.material, line.materiality_reason = _is_material(vals["month_var"], vals["month_budget"], client.materiality)
        lines.append(line)

    pack = Pack(client.client_id, client.name, client.currency, period.key, period.label, lines, checks,
                account_map, datetime.now(timezone.utc).isoformat(timespec="seconds"))
    _integrity_checks(pack, bundle, account_map)
    return pack


def _account_classes(bundle: SourceBundle) -> dict[str, str]:
    classes = {a.get("Code", ""): a.get("Class", "") for a in bundle.accounts}
    for line in bundle.tb_current + bundle.tb_prior:
        if line.code not in classes:
            classes[line.code] = "REVENUE" if line.section in ("Revenue", "Income") else \
                "EXPENSE" if line.section in ("Expenses", "Cost of Sales") else "BALANCE"
    return classes


def _integrity_checks(pack: Pack, bundle: SourceBundle, account_map: dict[str, str]) -> None:
    """Independent re-computation: the pack must agree with the trial balance it came from.

    Profit is recomputed from Xero's own account classes, not from our mapping, so a
    mapping error shows up as a mismatch instead of being silently agreed with.
    """
    classes = _account_classes(bundle)
    for code, line_key in account_map.items():
        is_pl_account = classes.get(code) in ("REVENUE", "EXPENSE")
        if is_pl_account != line_key.startswith("pl."):
            pack.checks.append(Check("MAP_WRONG_STATEMENT", "block",
                                     f"Account {code} ({classes.get(code)}) is mapped to {line_key}, which is on the other statement.",
                                     {"account": code, "line": line_key}))
    pl_lines = [l for l in bundle.tb_current if classes.get(l.code) in ("REVENUE", "EXPENSE")]
    tb_profit = -sum((l.ytd for l in pl_lines), ZERO)
    tb_month_profit = -sum((l.month for l in pl_lines), ZERO)
    net = pack.line("pl.net_profit").values
    if q2(net["ytd_actual"]) != q2(tb_profit):
        pack.checks.append(Check("PACK_YTD_PROFIT_MISMATCH", "block",
                                 f"Pack YTD profit {net['ytd_actual']:,.2f} differs from trial balance {tb_profit:,.2f}."))
    if q2(net["month_actual"]) != q2(tb_month_profit):
        pack.checks.append(Check("PACK_MONTH_PROFIT_MISMATCH", "block",
                                 f"Pack month profit {net['month_actual']:,.2f} differs from trial balance {tb_month_profit:,.2f}."))
    net_assets = pack.line("bs.net_assets").values["closing"]
    equity = pack.line("bs.total_equity").values["closing"]
    if q2(net_assets) != q2(equity):
        pack.checks.append(Check("BS_DOES_NOT_BALANCE", "block",
                                 f"Balance sheet does not balance: net assets {net_assets:,.2f} vs equity {equity:,.2f}."))
    for code in bundle.budgets:
        if code not in account_map:
            pack.checks.append(Check("BUDGET_UNMAPPED", "block", f"Budget exists for unmapped account {code}."))
    prior_codes = {p.code for p in bundle.tb_prior}
    for line in bundle.tb_current:
        if line.code not in prior_codes and line.month != 0:
            pack.checks.append(Check("NEW_ACCOUNT", "warn",
                                     f"Account {line.code} {line.name} has activity this month but none last month; check its mapping and coding.",
                                     {"account": line.code}))


# --------------------------------------------------------------------------- Excel

def _column_keys(ws) -> dict[int, str]:
    """Data sheets mark row 3 with '#keys' in column A; the rest of row 3 names the columns."""
    if ws.cell(row=3, column=1).value != "#keys":
        return {}
    return {cell.column: cell.value for cell in ws[3] if isinstance(cell.value, str) and cell.column > 1}


def _cell_value(value: Decimal | None, column: str):
    if value is None:
        return None
    if column.endswith("_pct"):
        return float(value) / 100  # Excel percent format
    return float(q2(value))


def write_pack_xlsx(pack: Pack, template: Path, out: Path, *, commentary: dict | None = None, meta: dict | None = None) -> None:
    from openpyxl import load_workbook

    wb = load_workbook(template)
    tokens = {
        "{{client_name}}": pack.client_name,
        "{{period_label}}": pack.period_label,
        "{{currency}}": pack.currency,
        "{{run_id}}": (meta or {}).get("run_id", ""),
        "{{prepared_at}}": (meta or {}).get("prepared_at", pack.generated_at),
    }
    by_key = {line.key: line for line in pack.lines}
    filled: set[str] = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    for token, value in tokens.items():
                        cell.value = cell.value.replace(token, str(value))
        col_keys = _column_keys(ws)
        if not col_keys:
            continue
        for row_idx in range(4, ws.max_row + 1):
            key = ws.cell(row=row_idx, column=1).value
            if key not in by_key:
                continue
            line = by_key[key]
            for col_idx, col_key in col_keys.items():
                if col_key in line.values:
                    ws.cell(row=row_idx, column=col_idx).value = _cell_value(line.values[col_key], col_key)
            filled.add(key)

    missing = [line.key for line in pack.lines if line.key not in filled]
    if missing:
        raise ValueError(f"template {template.name} has no row for: {', '.join(missing)}")

    if commentary is not None and "Commentary" in wb.sheetnames:
        ws = wb["Commentary"]
        row_idx = 5
        ws.cell(row=4, column=2).value = commentary.get("summary", "")
        for item in commentary.get("items", []):
            line = by_key.get(item["line_key"])
            if not line:
                continue
            ws.cell(row=row_idx + 1, column=1).value = line.key
            ws.cell(row=row_idx + 1, column=2).value = line.label
            ws.cell(row=row_idx + 1, column=3).value = _cell_value(line.values["month_actual"], "")
            ws.cell(row=row_idx + 1, column=4).value = _cell_value(line.values["month_budget"], "")
            ws.cell(row=row_idx + 1, column=5).value = _cell_value(line.values["month_var"], "")
            ws.cell(row=row_idx + 1, column=6).value = item["text"]
            ws.cell(row=row_idx + 1, column=7).value = "; ".join(e for d in item.get("drivers", []) for e in d.get("evidence", []))
            row_idx += 1
    if "Exceptions" in wb.sheetnames:
        ws = wb["Exceptions"]
        for offset, check in enumerate(pack.checks):
            ws.cell(row=5 + offset, column=1).value = check.severity.upper()
            ws.cell(row=5 + offset, column=2).value = check.code
            ws.cell(row=5 + offset, column=3).value = check.message
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def read_pack_xlsx(path: Path) -> dict[str, dict[str, float | None]]:
    """Read values back by key. The evaluation harness grades the file, not our in-memory objects."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    out: dict[str, dict[str, float | None]] = {}
    for ws in wb.worksheets:
        col_keys = _column_keys(ws)
        if not col_keys:
            continue
        for row_idx in range(4, ws.max_row + 1):
            key = ws.cell(row=row_idx, column=1).value
            if isinstance(key, str) and "." in key:
                out[key] = {ck: ws.cell(row=row_idx, column=ci).value for ci, ck in col_keys.items()}
    return out
