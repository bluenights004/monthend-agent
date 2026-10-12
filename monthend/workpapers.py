"""Referenced workpapers built from source documents.

Each workpaper row carries one or more references (evidence ids, see evidence.py)
plus the exact text found at that reference and the SHA-256 prefix of the file it
came from. The workpaper is the same whether a person or an agent reads it; agents
add *preparer notes* (judgement), never figures.

Tickmarks:  X agreed to Xero   S agreed to source document   E agreed to subsequent
            evidence (dated after period end)   C cast / recalculated   ! exception
"""

from __future__ import annotations

import json
from dataclasses import asdict, dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from monthend.config import ClientConfig, Period, WorkpaperSpec
from monthend.evidence import EvidenceIndex
from monthend.extraction import parse_bank_statement, parse_invoice, read_document
from monthend.money import ZERO, dec, q2
from monthend.pack import Check, Pack
from monthend.snapshot import SourceBundle
from monthend.sources.xero_parse import invoice_date, parse_xero_date

TICKMARKS = {
    "X": "Agreed to Xero ledger / trial balance",
    "S": "Agreed to source document",
    "E": "Agreed to subsequent evidence dated after period end",
    "C": "Cast / recalculated",
    "!": "Exception - see note",
}


@dataclass
class Ref:
    evidence_id: str
    source_hash: str
    quote: str


@dataclass
class WPRow:
    ref_no: str
    description: str
    amount: Decimal | None
    kind: str = "item"        # balance | reconciling | item | total | difference
    refs: list[Ref] = field(default_factory=list)
    tick: str = ""
    note: str = ""


@dataclass
class Workpaper:
    code: str
    title: str
    account: str
    pack_line: str
    objective: str
    rows: list[WPRow]
    conclusion: str
    issues: list[Check]
    tie: dict
    preparer_notes: list[str] = field(default_factory=list)

    def to_json(self) -> dict:
        return json.loads(json.dumps(asdict(self), default=str))

    @classmethod
    def from_json(cls, data: dict) -> "Workpaper":
        rows = [WPRow(**{**r, "amount": dec(r["amount"]) if r["amount"] is not None else None,
                         "refs": [Ref(**ref) for ref in r["refs"]]}) for r in data["rows"]]
        return cls(**{**data, "rows": rows, "issues": [Check(**c) for c in data["issues"]]})


class _Builder:
    def __init__(self, spec: WorkpaperSpec):
        self.spec, self.rows, self.issues = spec, [], []

    def row(self, description: str, amount: Decimal | None, kind: str = "item", refs=(), tick: str = "", note: str = "") -> WPRow:
        row = WPRow(f"{self.spec.code}.{len(self.rows) + 1}", description, q2(amount) if amount is not None else None,
                    kind, list(refs), tick, note)
        self.rows.append(row)
        return row

    def issue(self, code: str, severity: str, message: str, **detail) -> None:
        self.issues.append(Check(code, severity, f"{self.spec.code}: {message}", {"workpaper": self.spec.code, **detail}))


def _tb_ref(index: EvidenceIndex, code: str) -> Ref:
    ev = index.resolve(f"xero:tb:{code}")
    return Ref(f"xero:tb:{code}", ev.source_hash if ev else "", ev.text if ev else "not in trial balance")


def _doc_ref(index: EvidenceIndex, doc_id: str, locator: str) -> Ref:
    evidence_id = f"doc:{doc_id}#{locator}"
    ev = index.resolve(evidence_id)
    return Ref(evidence_id, ev.source_hash if ev else "", ev.text if ev else "unresolved")


def _tb_closing(bundle: SourceBundle, code: str) -> Decimal | None:
    line = next((l for l in bundle.tb_current if l.code == code), None)
    return line.ytd if line else None


# --------------------------------------------------------------------------- A: bank reconciliation

def bank_reconciliation(spec: WorkpaperSpec, client: ClientConfig, period: Period, bundle: SourceBundle,
                        index: EvidenceIndex) -> Workpaper:
    b = _Builder(spec)
    tolerance = dec(client.materiality.get("bank_rec_tolerance", "1.00"))
    ledger = _tb_closing(bundle, spec.account) or ZERO
    b.row(f"Balance per Xero, account {spec.account} at {period.end:%d %b %Y}", ledger, "balance",
          [_tb_ref(index, spec.account)], "X")

    statement_doc = bundle.document_by_filename(spec.options["statement"].format(period=period.key))
    if statement_doc is None:
        b.issue("WP_SOURCE_MISSING", "block", f"bank statement {spec.options['statement'].format(period=period.key)} not found in the client folder")
        return _finish(b, spec, "Bank statement not available; reconciliation not performed.", {"tb": ledger, "workpaper": None})
    statement = parse_bank_statement(statement_doc)
    if statement.closing is None:
        b.issue("WP_EXTRACTION_FAILED", "block", f"could not read the closing balance from {statement_doc.filename}")
        return _finish(b, spec, "Statement closing balance not readable.", {"tb": ledger, "workpaper": None})

    adjusted = ledger
    xero_amounts = []
    for txn in bundle.bank_transactions:
        if txn.get("BankAccount", {}).get("Code") != spec.account:
            continue
        when = parse_xero_date(txn.get("DateString") or txn.get("Date"))
        total = dec(txn.get("Total"))
        xero_amounts.append(total)
        if txn.get("IsReconciled") or (when and when > period.end):
            continue
        ref_name = txn.get("Reference") or txn["BankTransactionID"]
        ev_id = f"xero:banktxn:{ref_name}"
        ev = index.resolve(ev_id)
        refs = [Ref(ev_id, ev.source_hash if ev else "", ev.text if ev else "")]
        if txn.get("Type") == "SPEND":
            adjusted += total
            b.row(f"Add: unpresented payment - {txn['Contact']['Name']} {ref_name} ({when:%d %b})", total, "reconciling", refs, "X",
                  "Recorded in Xero, not yet cleared by the bank at period end. Confirm it clears on the October statement.")
        else:
            adjusted -= total
            b.row(f"Less: outstanding deposit - {txn['Contact']['Name']} {ref_name} ({when:%d %b})", -total, "reconciling", refs, "X",
                  "Recorded in Xero, banked after period end. Confirm it appears on the October statement.")

    unrecorded = []
    for line in statement.lines:
        amount = line.debit or line.credit
        if any(abs(amount - x) < Decimal("0.005") for x in xero_amounts):
            continue
        signed = -line.debit if line.debit else line.credit
        adjusted += signed
        unrecorded.append(line)
        b.row(f"{'Less' if signed < 0 else 'Add'}: on statement, not in Xero - {line.description} ({line.date})", signed,
              "reconciling", [_doc_ref(index, statement_doc.doc_id, line.locator)], "S!",
              "Not recorded in Xero. Proposed adjustment for the client to post (not posted by this system).")
    if unrecorded:
        total_unrecorded = sum((l.debit or l.credit for l in unrecorded), ZERO)
        b.issue("BANKREC_UNRECORDED_ITEMS", "warn",
                f"{len(unrecorded)} statement item(s) totalling {total_unrecorded:,.2f} are not recorded in Xero; proposed adjustment listed.",
                amount=str(total_unrecorded))

    b.row("Adjusted balance per Xero", adjusted, "total", tick="C")
    b.row(f"Balance per bank statement at {period.end:%d %b %Y}", statement.closing.value, "balance",
          [_doc_ref(index, statement_doc.doc_id, statement.closing.locator)], "S")
    difference = q2(adjusted - statement.closing.value)
    b.row("Unexplained difference", difference, "difference", tick="C" if abs(difference) <= tolerance else "C!")
    if abs(difference) > tolerance:
        b.issue("BANKREC_UNEXPLAINED", "block", f"unexplained bank reconciliation difference of {difference:,.2f}",
                difference=str(difference))
        conclusion = f"NOT RECONCILED: unexplained difference of {difference:,.2f} must be resolved before the pack is reviewed."
    else:
        conclusion = "Reconciled. Reconciling items are listed above and should clear in the following month."
    return _finish(b, spec, conclusion, {"tb": ledger, "workpaper": ledger})


# --------------------------------------------------------------------------- B: accruals

def accruals(spec: WorkpaperSpec, client: ClientConfig, period: Period, bundle: SourceBundle,
             index: EvidenceIndex) -> Workpaper:
    b = _Builder(spec)
    tolerance = dec(client.materiality.get("accrual_support_tolerance", "100"))
    ledger = -(_tb_closing(bundle, spec.account) or ZERO)  # credit balance shown positive
    schedule_name = spec.options["schedule"].format(period=period.key)
    schedule = bundle.document_by_filename(schedule_name)
    if schedule is None:
        b.issue("WP_SOURCE_MISSING", "block", f"accruals schedule {schedule_name} not found in the client folder")
        return _finish(b, spec, "Accruals schedule not available.", {"tb": ledger, "workpaper": None})

    total = ZERO
    for n, row in enumerate(read_document(schedule).rows, start=1):
        amount = dec(row["amount_ex_gst"])
        total += amount
        refs = [_doc_ref(index, schedule.doc_id, f"r{n}:amount_ex_gst")]
        tick, notes = "S", []
        support_id = Path(row.get("support", "")).stem
        support = bundle.documents.get(support_id)
        if not row.get("support") or support is None:
            tick += "!"
            notes.append("No supporting document found in the client folder.")
            b.issue("ACCRUAL_NO_SUPPORT", "warn", f"{row['ref']} {row['supplier']} {amount:,.2f} has no supporting document",
                    accrual=row["ref"])
        else:
            invoice = parse_invoice(support)
            if invoice.subtotal is None:
                tick += "!"
                notes.append(f"Could not read the subtotal from {support.filename}.")
                b.issue("WP_EXTRACTION_FAILED", "warn", f"could not read subtotal from {support.filename}")
            else:
                refs.append(_doc_ref(index, support.doc_id, invoice.subtotal.locator))
                inv_date = None
                if invoice.invoice_date:
                    day, month, year = invoice.invoice_date.text.split("/")
                    inv_date = date(int(year), int(month), int(day))
                    refs.append(_doc_ref(index, support.doc_id, invoice.invoice_date.locator))
                if inv_date and inv_date > period.end:
                    tick += "E"
                diff = q2(invoice.subtotal.value - amount)
                if not invoice.arithmetic_ok():
                    tick += "!"
                    notes.append("Invoice subtotal + GST does not equal total; check the document.")
                    b.issue("SUPPORT_ARITHMETIC", "warn", f"{support.filename} does not add up")
                if abs(diff) > tolerance:
                    tick += "!"
                    notes.append(f"Invoice {invoice.invoice_number.text if invoice.invoice_number else ''} is {diff:,.2f} different from the accrual.")
                    b.issue("ACCRUAL_SUPPORT_DIFFERENCE", "warn",
                            f"{row['ref']} accrual {amount:,.2f} vs invoice {invoice.subtotal.value:,.2f} (difference {diff:,.2f})",
                            accrual=row["ref"], difference=str(diff))
                elif diff != 0:
                    notes.append(f"Invoice is {diff:,.2f} different from the accrual; below the {tolerance:,.0f} tolerance, no adjustment proposed.")
                else:
                    notes.append(f"Agreed to invoice {invoice.invoice_number.text if invoice.invoice_number else support.filename}.")
        b.row(f"{row['ref']} {row['supplier']} - {row['description']} (account {row['account_code']})", amount, "item",
              refs, tick, " ".join(notes))

    b.row("Total per accruals schedule", total, "total", tick="C")
    b.row(f"Balance per Xero, account {spec.account}", ledger, "balance", [_tb_ref(index, spec.account)], "X")
    difference = q2(total - ledger)
    b.row("Difference", difference, "difference", tick="C!" if abs(difference) > Decimal("1") else "C")
    if abs(difference) > Decimal("1"):
        b.issue("ACCRUALS_DO_NOT_TIE", "block", f"accruals schedule differs from the ledger by {difference:,.2f}",
                difference=str(difference))
        conclusion = f"NOT AGREED: schedule differs from the ledger by {difference:,.2f}."
    else:
        conclusion = "Accrued expenses agree to the schedule; each accrual was checked against supporting invoices where available."
    return _finish(b, spec, conclusion, {"tb": ledger, "workpaper": total})


# --------------------------------------------------------------------------- C: receivables

def _outstanding_at(invoice: dict, as_at: date) -> Decimal:
    payments = invoice.get("Payments")
    if payments is None:
        return dec(invoice.get("AmountDue"))
    paid = sum((dec(p.get("Amount")) for p in payments
                if (parse_xero_date(p.get("Date") or p.get("DateString")) or as_at) <= as_at), ZERO)
    return dec(invoice.get("Total")) - paid


def receivables(spec: WorkpaperSpec, client: ClientConfig, period: Period, bundle: SourceBundle,
                index: EvidenceIndex) -> Workpaper:
    b = _Builder(spec)
    overdue_days = int(client.materiality.get("receivables_overdue_days", 60))
    ledger = _tb_closing(bundle, spec.account) or ZERO
    buckets = {"0-30 days": ZERO, "31-60 days": ZERO, "61-90 days": ZERO, "Over 90 days": ZERO}
    total = ZERO
    open_items = []
    for inv in bundle.invoices:
        when = invoice_date(inv)
        if inv.get("Type") != "ACCREC" or when is None or when > period.end:
            continue
        outstanding = _outstanding_at(inv, period.end)
        if outstanding <= 0:
            continue
        open_items.append((when, inv, outstanding))
    for when, inv, outstanding in sorted(open_items, key=lambda item: item[0]):
        age = (period.end - when).days
        bucket = "0-30 days" if age <= 30 else "31-60 days" if age <= 60 else "61-90 days" if age <= 90 else "Over 90 days"
        buckets[bucket] += outstanding
        total += outstanding
        ev_id = f"xero:invoice:{inv['InvoiceNumber']}"
        ev = index.resolve(ev_id)
        tick, note = "X", f"Invoice dated {when:%d %b %Y}, {age} days old at period end ({bucket})."
        if age >= overdue_days:
            tick += "!"
            note += " Aged: confirm recoverability with the client; consider whether a provision is needed."
            b.issue("RECEIVABLE_AGED", "warn", f"{inv['InvoiceNumber']} {inv['Contact']['Name']} {outstanding:,.2f} is {age} days old",
                    invoice=inv["InvoiceNumber"], amount=str(outstanding))
        b.row(f"{inv['InvoiceNumber']} {inv['Contact']['Name']}", outstanding, "item",
              [Ref(ev_id, ev.source_hash if ev else "", ev.text if ev else "")], tick, note)
    b.row("Total open sales invoices", total, "total", tick="C")
    for name, amount in buckets.items():
        b.row(f"  of which {name}", amount, "analysis", tick="C")
    b.row(f"Balance per Xero, account {spec.account}", ledger, "balance", [_tb_ref(index, spec.account)], "X")
    difference = q2(total - ledger)
    b.row("Difference", difference, "difference", tick="C!" if abs(difference) > Decimal("1") else "C")
    if abs(difference) > Decimal("1"):
        b.issue("RECEIVABLES_DO_NOT_TIE", "block", f"open invoices differ from the receivables ledger by {difference:,.2f}",
                difference=str(difference))
        conclusion = f"NOT AGREED: open invoices differ from the ledger by {difference:,.2f}."
    else:
        conclusion = "Trade receivables agree to open sales invoices. Aged items are flagged for follow-up."
    return _finish(b, spec, conclusion, {"tb": ledger, "workpaper": total})


# --------------------------------------------------------------------------- common

def _finish(b: _Builder, spec: WorkpaperSpec, conclusion: str, tie: dict) -> Workpaper:
    objectives = {
        "bank_reconciliation": "Agree the Xero cash balance to the bank statement and explain every reconciling item.",
        "accruals": "Agree accrued expenses to the schedule and test each accrual against supporting invoices.",
        "receivables": "Agree trade receivables to open sales invoices and identify aged balances.",
    }
    return Workpaper(spec.code, spec.title, spec.account, spec.pack_line, objectives.get(spec.kind, ""), b.rows,
                     conclusion, b.issues, {k: (str(q2(v)) if isinstance(v, Decimal) else v) for k, v in tie.items()})


BUILDERS = {"bank_reconciliation": bank_reconciliation, "accruals": accruals, "receivables": receivables}


def build_workpapers(client: ClientConfig, period: Period, bundle: SourceBundle, index: EvidenceIndex) -> list[Workpaper]:
    papers = []
    for spec in client.workpapers:
        builder = BUILDERS.get(spec.kind)
        if builder is None:
            raise ValueError(f"no builder for workpaper kind {spec.kind!r}")
        paper = builder(spec, client, period, bundle, index)
        for row in paper.rows:
            index.add_workpaper_row(paper.code, row.ref_no, row.description, row.amount)
        papers.append(paper)
    return papers


def lead_schedule(pack: Pack, papers: list[Workpaper]) -> list[dict]:
    rows = []
    for paper in papers:
        line = pack.line(paper.pack_line)
        rows.append({"workpaper": paper.code, "pack_line": line.key, "label": line.label,
                     "pack_closing": str(q2(line.values["closing"])), "accounts": ", ".join(line.accounts),
                     "workpaper_account": paper.account, "workpaper_balance": paper.tie.get("workpaper")})
    return rows


# --------------------------------------------------------------------------- outputs

def write_workpapers_xlsx(papers: list[Workpaper], pack: Pack, bundle: SourceBundle, out: Path, run_id: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    head_fill, head_font = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF")
    wb = Workbook()
    idx = wb.active
    idx.title = "Index"
    idx["A1"], idx["A1"].font = f"{pack.client_name} - workpapers {pack.period_label}", Font(bold=True, size=14)
    idx["A2"] = f"Prepared by monthend-agent, run {run_id}. Reviewed by: ________  Date: ________"
    for i, label in enumerate(["WP", "Title", "Pack line", "Ledger", "Workpaper", "Status", "Issues"], start=1):
        c = idx.cell(row=4, column=i, value=label)
        c.fill, c.font = head_fill, head_font
    for r, paper in enumerate(papers, start=5):
        status = "Exceptions" if any(i.severity == "block" for i in paper.issues) else "Notes" if paper.issues else "Clear"
        for c, value in enumerate([paper.code, paper.title, paper.pack_line, paper.tie.get("tb"), paper.tie.get("workpaper"),
                                   status, len(paper.issues)], start=1):
            idx.cell(row=r, column=c, value=float(value) if c in (4, 5) and value is not None else value)
    legend_row = 6 + len(papers)
    idx.cell(row=legend_row, column=1, value="Tickmarks").font = Font(bold=True)
    for offset, (mark, meaning) in enumerate(TICKMARKS.items(), start=1):
        idx.cell(row=legend_row + offset, column=1, value=mark)
        idx.cell(row=legend_row + offset, column=2, value=meaning)
    for col, width in zip("ABCDEFG", (8, 48, 18, 14, 14, 12, 8)):
        idx.column_dimensions[col].width = width

    for paper in papers:
        ws = wb.create_sheet(paper.code)
        ws["A1"], ws["A1"].font = f"{paper.code} {paper.title}", Font(bold=True, size=13)
        ws["A2"] = f"Objective: {paper.objective}"
        ws["A3"] = f"Pack line: {paper.pack_line}   Account: {paper.account}   Run: {run_id}"
        for i, label in enumerate(["Ref", "Description", "Amount", "Tick", "Source reference", "Source hash", "Text at source", "Note"], start=1):
            c = ws.cell(row=5, column=i, value=label)
            c.fill, c.font = head_fill, head_font
        r = 6
        for row in paper.rows:
            refs = row.refs or [Ref("", "", "")]
            for n, ref in enumerate(refs):
                ws.cell(row=r, column=1, value=row.ref_no if n == 0 else "")
                ws.cell(row=r, column=2, value=row.description if n == 0 else "")
                if n == 0 and row.amount is not None:
                    ws.cell(row=r, column=3, value=float(row.amount)).number_format = '#,##0.00;(#,##0.00);"-"'
                ws.cell(row=r, column=4, value=row.tick if n == 0 else "")
                ws.cell(row=r, column=5, value=ref.evidence_id)
                ws.cell(row=r, column=6, value=ref.source_hash)
                ws.cell(row=r, column=7, value=ref.quote)
                ws.cell(row=r, column=8, value=row.note if n == 0 else "").alignment = Alignment(wrap_text=True, vertical="top")
                if row.kind in ("total", "balance", "difference") and n == 0:
                    for c in (2, 3):
                        ws.cell(row=r, column=c).font = Font(bold=True)
                r += 1
        r += 1
        ws.cell(row=r, column=1, value="Conclusion").font = Font(bold=True)
        ws.cell(row=r, column=2, value=paper.conclusion)
        for issue in paper.issues:
            r += 1
            ws.cell(row=r, column=1, value=issue.severity.upper())
            ws.cell(row=r, column=2, value=issue.message)
        if paper.preparer_notes:
            r += 2
            ws.cell(row=r, column=1, value="Preparer notes (agent)").font = Font(bold=True)
            for note in paper.preparer_notes:
                r += 1
                ws.cell(row=r, column=2, value=note).alignment = Alignment(wrap_text=True)
        for col, width in zip("ABCDEFGH", (8, 60, 14, 6, 44, 14, 60, 60)):
            ws.column_dimensions[col].width = width

    lead = wb.create_sheet("Lead schedule")
    for i, label in enumerate(["WP", "Pack line", "Label", "Pack balance", "Accounts in line", "WP account", "WP balance"], start=1):
        c = lead.cell(row=1, column=i, value=label)
        c.fill, c.font = head_fill, head_font
    for r, row in enumerate(lead_schedule(pack, papers), start=2):
        for c, key in enumerate(["workpaper", "pack_line", "label", "pack_closing", "accounts", "workpaper_account", "workpaper_balance"], start=1):
            value = row[key]
            lead.cell(row=r, column=c, value=float(value) if key in ("pack_closing", "workpaper_balance") and value else value)

    src = wb.create_sheet("Sources")
    for i, label in enumerate(["Document id", "File", "SHA-256", "Type"], start=1):
        c = src.cell(row=1, column=i, value=label)
        c.fill, c.font = head_fill, head_font
    for r, doc in enumerate(sorted(bundle.documents.values(), key=lambda d: d.doc_id), start=2):
        for c, value in enumerate([doc.doc_id, doc.filename, doc.sha256, doc.media_type], start=1):
            src.cell(row=r, column=c, value=value)
    src.column_dimensions["A"].width, src.column_dimensions["C"].width = 36, 68
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def workpapers_markdown(papers: list[Workpaper], pack: Pack, run_id: str) -> str:
    out = [f"# Workpapers - {pack.client_name}, {pack.period_label}", "", f"Run `{run_id}`. Figures in {pack.currency}.", ""]
    for paper in papers:
        out += [f"## {paper.code} {paper.title}", "", f"*Objective:* {paper.objective}", "",
                "| Ref | Description | Amount | Tick | Source |", "|---|---|---:|---|---|"]
        for row in paper.rows:
            amount = f"{row.amount:,.2f}" if row.amount is not None else ""
            sources = "<br>".join(f"`{r.evidence_id}` [{r.source_hash}]" for r in row.refs)
            out.append(f"| {row.ref_no} | {row.description}{(' - ' + row.note) if row.note else ''} | {amount} | {row.tick} | {sources} |")
        out += ["", f"**Conclusion:** {paper.conclusion}", ""]
        for issue in paper.issues:
            out.append(f"- **{issue.severity.upper()}** {issue.message}")
        if paper.preparer_notes:
            out += ["", "*Preparer notes (agent):*"] + [f"- {n}" for n in paper.preparer_notes]
        out.append("")
    out += ["Tickmarks: " + "; ".join(f"{k} {v}" for k, v in TICKMARKS.items()), ""]
    return "\n".join(out)
