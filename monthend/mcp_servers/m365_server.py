"""MCP server: Microsoft 365 (Outlook, SharePoint, Excel) through Microsoft Graph.

Read tools are open. The single write tool, `deliver_approved_pack`, takes only a
run id: recipients, subject and attachments come from the approved manifest, and the
review gate refuses unless a person approved that exact bundle.
"""

from __future__ import annotations

import json
from datetime import datetime

from mcp.server.fastmcp import FastMCP

from monthend.config import Period, load_client
from monthend.mcp_servers._common import tool_errors

mcp = FastMCP("m365", instructions="Outlook, SharePoint and Excel. Cite emails as email:<id>. You cannot send anything that a person has not approved.")


def _live(client) -> bool:
    return client.backend == "live"


def _graph():
    from monthend.sources.graph import GraphClient

    return GraphClient.from_env()


def _messages(client) -> list[dict]:
    return json.loads((client.fixtures_dir / "outlook" / "messages.json").read_text())["value"]


@mcp.tool()
@tool_errors
def outlook_search_messages(client_id: str, received_after: str, received_before: str, text: str = "") -> list[dict]:
    """Messages in the client mailbox received between two ISO dates, optionally containing text. Returns previews only."""
    client = load_client(client_id)
    if _live(client):
        messages = _graph().list_messages(client.mailbox, f"{received_after}T00:00:00Z", f"{received_before}T23:59:59Z")
    else:
        messages = [m for m in _messages(client) if received_after <= m["receivedDateTime"][:10] <= received_before]
    needle = text.lower()
    out = []
    for m in messages:
        blob = f"{m.get('subject', '')} {m.get('bodyPreview', '')}".lower()
        if needle and needle not in blob:
            continue
        out.append({"evidence_id": f"email:{m['id']}", "received": m["receivedDateTime"], "subject": m.get("subject"),
                    "from": m.get("from", {}).get("emailAddress", {}).get("address"), "preview": m.get("bodyPreview")})
    return out


@mcp.tool()
@tool_errors
def sharepoint_list_client_documents(client_id: str, period: str) -> list[dict]:
    """Files in the client's SharePoint folder for a period (YYYY-MM)."""
    client = load_client(client_id)
    folder = client.sharepoint_folder.format(period=Period.parse(period).key)
    if _live(client):
        from monthend.sources.graph import sharepoint_site

        graph = _graph()
        site = graph.site_id(*sharepoint_site())
        return [{"name": i["name"], "size": i.get("size"), "modified": i.get("lastModifiedDateTime")}
                for i in graph.folder_children(site, folder) if "file" in i]
    path = client.fixtures_dir / "sharepoint" / folder
    return [{"name": p.name, "size": p.stat().st_size, "modified": datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec="seconds")}
            for p in sorted(path.iterdir()) if p.is_file()] if path.exists() else []


@mcp.tool()
@tool_errors
def excel_read_range(client_id: str, period: str, filename: str, sheet: str, address: str) -> list[list]:
    """Read a cell range (e.g. 'A1:F20') from an Excel workbook in the client's period folder."""
    client = load_client(client_id)
    folder = client.sharepoint_folder.format(period=Period.parse(period).key)
    if _live(client):
        from monthend.sources.graph import sharepoint_site

        graph = _graph()
        site = graph.site_id(*sharepoint_site())
        item = next(i for i in graph.folder_children(site, folder) if i["name"] == filename)
        return graph.read_range(site, item["id"], sheet, address)
    from openpyxl import load_workbook

    ws = load_workbook(client.fixtures_dir / "sharepoint" / folder / filename, data_only=True, read_only=True)[sheet]
    return [[cell.value for cell in row] for row in ws[address]]


@mcp.tool()
@tool_errors
def deliver_approved_pack(run_id: str) -> dict:
    """Deliver an APPROVED run: Outlook draft (or send, if the client is configured to) and SharePoint upload.

    Refused unless a person approved this exact bundle and nothing has changed since.
    """
    from monthend import pipeline

    return {"released": pipeline.release(run_id, actor="agent:release-after-human-approval")}


def main() -> None:
    mcp.run()


if __name__ == "__main__":
    main()
