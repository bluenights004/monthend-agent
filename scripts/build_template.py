"""Generate templates/management-pack.xlsx from a mapping file.

Firms normally bring their own template; this script produces a clean starting
point that follows the key conventions in docs/07-customising-the-template.md:

- column A (hidden) holds the row key, e.g. pl.rev.contract
- row 3 (hidden) starts with '#keys' and names each data column, e.g. month_actual
- {{client_name}}, {{period_label}}, {{currency}}, {{run_id}}, {{prepared_at}} tokens are replaced

    python scripts/build_template.py [config/mappings/standard-au.json]
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
NUM = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
TOTAL_BORDER = Border(top=Side(style="thin"), bottom=Side(style="double"))
TITLE = Font(bold=True, size=14, color="1F3A5F")

PL_COLS = [("month_actual", "Actual", NUM), ("month_budget", "Budget", NUM), ("month_var", "Variance", NUM),
           ("month_var_pct", "Var %", PCT), ("prior_month", "Prior month", NUM),
           ("ytd_actual", "YTD actual", NUM), ("ytd_budget", "YTD budget", NUM), ("ytd_var", "YTD variance", NUM)]
BS_COLS = [("closing", "This month", NUM), ("prior_closing", "Prior month", NUM), ("movement", "Movement", NUM)]


def data_sheet(wb, title, subtitle, rows, cols):
    ws = wb.create_sheet(title)
    ws["B1"] = "{{client_name}}"
    ws["B1"].font = TITLE
    ws["B2"] = subtitle
    ws["A3"] = "#keys"
    ws.cell(row=4, column=2, value="").fill = HEAD_FILL
    for i, (key, label, _) in enumerate(cols, start=3):
        ws.cell(row=3, column=i, value=key)
        head = ws.cell(row=4, column=i, value=label)
        head.fill, head.font = HEAD_FILL, HEAD_FONT
        head.alignment = Alignment(horizontal="right")
        ws.column_dimensions[get_column_letter(i)].width = 14
    positions = {}
    r = 5
    for row in rows:
        ws.cell(row=r, column=1, value=row["key"])
        label = ws.cell(row=r, column=2, value=row["label"])
        is_total = row["type"] in ("total",)
        for i, (_, _, fmt) in enumerate(cols, start=3):
            cell = ws.cell(row=r, column=i)
            cell.number_format = fmt
            if is_total:
                cell.font, cell.border = Font(bold=True), TOTAL_BORDER
        if is_total:
            label.font = Font(bold=True)
        positions[row["key"]] = r
        r += 1
        if is_total:
            r += 1  # spacer row after totals
    ws.column_dimensions["A"].hidden = True
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["B"].width = 42
    ws.freeze_panes = "C5"
    return ws, positions


def main() -> None:
    mapping_path = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "config" / "mappings" / "standard-au.json"
    rows = json.loads(mapping_path.read_text())["rows"]
    wb = Workbook()
    cover = wb.active
    cover.title = "Cover"
    cover["B2"] = "{{client_name}}"
    cover["B2"].font = Font(bold=True, size=20, color="1F3A5F")
    cover["B4"] = "Management accounts for {{period_label}}"
    cover["B4"].font = Font(size=14)
    cover["B6"] = "Currency: {{currency}}"
    cover["B7"] = "Prepared: {{prepared_at}}"
    cover["B8"] = "Preparation reference: {{run_id}}"
    cover["B10"] = "Contents: Profit and loss, Balance sheet, Commentary, Exceptions, Checks"
    cover["B12"] = ("Basis of preparation: unaudited management accounts prepared from the Xero general ledger. "
                    "Budget figures are the board-approved budget held in Xero.")
    cover["B12"].alignment = Alignment(wrap_text=True)
    cover.column_dimensions["B"].width = 100

    pl_rows = [r for r in rows if r["statement"] == "pl"]
    bs_rows = [r for r in rows if r["statement"] == "bs"]
    pl, pl_pos = data_sheet(wb, "P&L", "Profit and loss - {{period_label}} ({{currency}})", pl_rows, PL_COLS)
    bs, bs_pos = data_sheet(wb, "Balance Sheet", "Balance sheet as at end of {{period_label}} ({{currency}})", bs_rows, BS_COLS)

    com = wb.create_sheet("Commentary")
    com["B1"] = "Commentary on material budget variances"
    com["B1"].font = TITLE
    com["A4"] = "Summary"
    com["B4"].alignment = Alignment(wrap_text=True, vertical="top")
    for i, label in enumerate(["Key", "Line", "Actual", "Budget", "Variance", "Commentary", "Evidence"], start=1):
        cell = com.cell(row=5, column=i, value=label)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    for col, width in zip("ABCDEFG", (18, 34, 13, 13, 13, 90, 40)):
        com.column_dimensions[col].width = width
    for r in range(6, 40):
        for c in (3, 4, 5):
            com.cell(row=r, column=c).number_format = NUM
        com.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")

    exc = wb.create_sheet("Exceptions")
    exc["B1"] = "Exceptions raised during preparation"
    exc["B1"].font = TITLE
    for i, label in enumerate(["Severity", "Code", "Message"], start=1):
        cell = exc.cell(row=4, column=i, value=label)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    exc.column_dimensions["B"].width, exc.column_dimensions["C"].width = 26, 110

    chk = wb.create_sheet("Checks")
    chk["B1"] = "Arithmetic checks (live Excel formulas; every difference should be 0)"
    chk["B1"].font = TITLE
    for i, label in enumerate(["Check", "Difference"], start=2):
        cell = chk.cell(row=4, column=i, value=label)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
    r = 5
    for sheet, positions, sheet_rows, col in (("P&L", pl_pos, pl_rows, "C"), ("P&L", pl_pos, pl_rows, "H"),
                                               ("Balance Sheet", bs_pos, bs_rows, "C")):
        for row in sheet_rows:
            if row["type"] != "total":
                continue
            terms = "".join(("-" if t.startswith("-") else "+") + f"'{sheet}'!{col}{positions[t.lstrip('-')]}" for t in row["sum"])
            chk.cell(row=r, column=2, value=f"{sheet} {row['label']} (column {col})")
            cell = chk.cell(row=r, column=3, value=f"=ROUND('{sheet}'!{col}{positions[row['key']]}-({terms}),2)")
            cell.number_format = "0.00"
            r += 1
    chk.cell(row=r, column=2, value="Balance sheet balances (net assets - equity)")
    chk.cell(row=r, column=3, value=f"=ROUND('Balance Sheet'!C{bs_pos['bs.net_assets']}-'Balance Sheet'!C{bs_pos['bs.total_equity']},2)")
    chk.cell(row=r + 1, column=2, value="P&L net profit equals current year earnings (YTD)")
    chk.cell(row=r + 1, column=3, value=f"=ROUND('P&L'!H{pl_pos['pl.net_profit']}-'Balance Sheet'!C{bs_pos['bs.current_earnings']},2)")
    chk.column_dimensions["B"].width = 70

    out = ROOT / "templates" / "management-pack.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print(f"wrote {out.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
